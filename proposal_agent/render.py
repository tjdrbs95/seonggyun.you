"""제안서 덱을 PowerPoint 템플릿(기본: Workday Corporate Template) 위에 그린다.

슬라이드 배경·로고·제목 스타일은 템플릿 레이아웃을 그대로 쓰고, 본문 도형은
템플릿 예시 장표(TCO 장표)의 서식(맑은 고딕, 네이비 박스, 옅은 청록 테두리 박스)을 따른다.
"""

from __future__ import annotations

import copy
from dataclasses import dataclass
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, PP_PLACEHOLDER
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

from .deck import STATUS_LABELS, ProposalDeck, TableSlide
from .layout import (
    BULLET_INDENT_IN,
    CARD_PAD,
    MIN_CARD_H,
    CARD_HEADER_H,
    LINE_SPACING,
    SUB_INDENT_IN,
    TABLE_SIZES,
    Box,
    Geometry,
    card_inner,
    column_boxes,
    fitted_height,
    plan_slide,
    process_boxes,
    bullets_height as _bh,
    split_bullet,
    table_row_heights,
    table_widths,
)


@dataclass(frozen=True)
class Theme:
    font: str = "Malgun Gothic"
    navy: str = "0F2E66"  # 제목, 헤드 메시지
    text: str = "022043"  # 본문
    box: str = "1F497D"  # 네이비 박스(예시 장표의 Product/Deployment 박스)
    accent: str = "1C98E8"  # 강조(Workday 블루)
    accent_light: str = "9ECFFF"
    outline: str = "4BACC6"  # 둥근 테두리 박스(50% 투명)
    muted: str = "7F8B9B"
    row_alt: str = "F3F6FA"
    grid: str = "C9D3E0"
    footer: str = "Workday Confidential"


# 레이아웃 역할 → 템플릿 레이아웃 이름 후보(앞쪽 우선)
LAYOUT_CANDIDATES = {
    "cover": ["Title Slide"],
    "content": ["Title Only"],
    "headline": ["1/2_Headline", "Title Only"],
    "section": ["Section Title", "Section Header", "Title Only"],
    "closing": ["Bumper Slide", "Blank", "Title Only"],
}

APPENDIX_TITLE = "부록. RFP 요구사항 대응표"
APPENDIX_ROWS_MAX = 12  # TableSlide.rows 최대값과 같게


def _rgb(hex_: str) -> RGBColor:
    return RGBColor.from_string(hex_)


class DeckRenderer:
    def __init__(self, template: Path, theme: Theme | None = None) -> None:
        self.template = Path(template)
        prs = Presentation(str(self.template))
        self.theme = theme or theme_from_template(prs)
        self.geometry = Geometry(Emu(prs.slide_width).inches, Emu(prs.slide_height).inches)

    # ── 공개 API ────────────────────────────────────────────

    def render(self, deck: ProposalDeck, out_path: Path) -> list:
        """덱을 그려 저장하고, 실제로 그린 슬라이드 목록(부록 포함)을 돌려준다."""
        prs = Presentation(str(self.template))
        _clear_slides(prs)
        layouts = {name: _find_layout(prs, name) for name in LAYOUT_CANDIDATES}

        slides = with_appendix(deck, self.geometry)
        for number, spec in enumerate(slides, 1):
            role = {"cover": "cover", "section": "section", "closing": "closing", "headline": "headline"}.get(
                spec.kind, "content"
            )
            slide = prs.slides.add_slide(layouts[role])
            plan = plan_slide(spec, self.geometry)
            getattr(self, f"_draw_{spec.kind}")(slide, spec, plan)
            if role != "cover" and spec.kind != "closing":
                self._footer(slide, number, getattr(spec, "rfp_refs", []))
            _drop_empty_placeholders(slide)
            notes = getattr(spec, "notes", "")
            if notes:
                slide.notes_slide.notes_text_frame.text = notes

        out_path.parent.mkdir(parents=True, exist_ok=True)
        prs.save(str(out_path))
        return slides

    # ── 종류별 그리기 ─────────────────────────────────────────

    def _draw_cover(self, slide, spec, plan) -> None:
        title_pt = 28 if len(spec.title) <= 26 else 22
        self._set_placeholder(slide, PP_PLACEHOLDER.CENTER_TITLE, spec.title, title_pt)
        if spec.subtitle:
            self._set_placeholder(slide, PP_PLACEHOLDER.SUBTITLE, spec.subtitle, 14)

    def _draw_closing(self, slide, spec, plan) -> None:
        pass

    def _draw_section(self, slide, spec, plan) -> None:
        t = self.theme
        left = 3.2 if self.geometry.width >= 9 else 1.0
        width = self.geometry.width - left - 0.5
        mid = self.geometry.height / 2
        if spec.number:
            self._text(slide, Box(left, mid - 1.05, width, 0.6), spec.number, 32, t.accent, bold=True)
        self._text(slide, Box(left, mid - 0.45, width, 1.0), spec.title, plan.title_pt, t.navy, bold=True)
        if spec.subtitle:
            self._text(slide, Box(left, mid + 0.55, width, 0.6), spec.subtitle, 12, t.text)

    def _draw_agenda(self, slide, spec, plan) -> None:
        t = self.theme
        self._title(slide, spec.title, plan.title_pt)
        body = self.geometry.body_box(False)
        per_col, cols = plan.extra["per_col"], plan.extra["cols"]
        col_w = (body.width - 0.2 * (cols - 1)) / cols
        row_h = min(0.62, body.height / per_col)
        for i, item in enumerate(spec.items):
            c, r = divmod(i, per_col)
            x, y = body.left + c * (col_w + 0.2), body.top + r * row_h
            num = self._rect(slide, Box(x, y + 0.08, 0.5, row_h - 0.16), fill=t.box)
            self._fill_text(num.text_frame, f"{i + 1:02d}", 12, "FFFFFF", bold=True, align=PP_ALIGN.CENTER)
            self._text(slide, Box(x + 0.65, y + 0.08, col_w - 0.7, row_h - 0.16), item, plan.body_pt, t.navy,
                       bold=True, anchor=MSO_ANCHOR.MIDDLE)
            line = slide.shapes.add_connector(1, Inches(x), Inches(y + row_h), Inches(x + col_w), Inches(y + row_h))
            line.line.color.rgb = _rgb(t.grid)
            line.line.width = Pt(0.75)

    def _draw_bullets(self, slide, spec, plan) -> None:
        self._title_message(slide, spec, plan)
        body = self.geometry.body_box(bool(spec.message))
        self._bullets(slide, Box(body.left, body.top, body.width, body.height), spec.bullets, plan.body_pt)

    def _draw_two_column(self, slide, spec, plan) -> None:
        self._title_message(slide, spec, plan)
        self._cards(slide, [spec.left, spec.right], bool(spec.message), plan.body_pt)

    def _draw_cards(self, slide, spec, plan) -> None:
        self._title_message(slide, spec, plan)
        self._cards(slide, spec.cards, bool(spec.message), plan.body_pt)

    def _draw_table(self, slide, spec, plan) -> None:
        t = self.theme
        self._title_message(slide, spec, plan)
        body = self.geometry.body_box(bool(spec.message))
        widths = plan.extra["widths"]
        heights = table_row_heights([spec.columns] + spec.rows, widths, plan.body_pt)
        shape = slide.shapes.add_table(
            len(spec.rows) + 1, len(spec.columns), Inches(body.left), Inches(body.top),
            Inches(body.width), Inches(sum(heights)),
        )
        table = shape.table
        _plain_table_style(shape)
        for j, w in enumerate(widths):
            table.columns[j].width = Inches(w)
        for i, h in enumerate(heights):
            table.rows[i].height = Inches(h)
        for i, row in enumerate([spec.columns] + spec.rows):
            header = i == 0
            for j, value in enumerate(row):
                cell = table.cell(i, j)
                cell.margin_left = cell.margin_right = Inches(0.06)
                cell.margin_top = cell.margin_bottom = Inches(0.04)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                cell.fill.solid()
                cell.fill.fore_color.rgb = _rgb(t.box if header else ("FFFFFF" if i % 2 else t.row_alt))
                self._fill_text(
                    cell.text_frame, value, plan.body_pt,
                    "FFFFFF" if header else t.text, bold=header,
                    align=PP_ALIGN.CENTER if header else PP_ALIGN.LEFT,
                )
                _cell_borders(cell, t.grid)

    def _draw_process(self, slide, spec, plan) -> None:
        t = self.theme
        self._title_message(slide, spec, plan)
        chevrons, details = process_boxes(self.geometry, bool(spec.message), len(spec.steps))
        detail_h = max(
            [min(d.height, max(MIN_CARD_H, 2 * CARD_PAD + 0.1 + _bh(step.bullets, d.width - 0.16, plan.body_pt)))
             for step, d in zip(spec.steps, details) if step.bullets] or [MIN_CARD_H]
        )
        for i, (step, c, d) in enumerate(zip(spec.steps, chevrons, details)):
            d = Box(d.left, d.top, d.width, detail_h)
            shape_type = MSO_SHAPE.PENTAGON if i == 0 else MSO_SHAPE.CHEVRON
            chev = self._rect(slide, c, fill=t.box, shape=shape_type)
            tf = chev.text_frame
            tf.margin_left = Inches(0.22 if i else 0.08)
            tf.margin_right = Inches(0.12)
            self._fill_text(tf, step.name, plan.extra["step_pt"][i], "FFFFFF", bold=True, align=PP_ALIGN.CENTER)
            if step.period:
                p = tf.add_paragraph()
                p.alignment = PP_ALIGN.CENTER
                self._run(p, step.period, 8, t.accent_light)
            if step.bullets:
                box = self._rect(slide, d, fill=None, line=t.outline, rounded=True)
                self._bullets_in(box.text_frame, step.bullets, plan.body_pt, margin=0.08)

    def _draw_headline(self, slide, spec, plan) -> None:
        t = self.theme
        if not self._set_placeholder(slide, PP_PLACEHOLDER.TITLE, spec.title, plan.title_pt):
            self._text(slide, Box(0.5, 0.41, 3.7, 2.2), spec.title, plan.title_pt, t.navy, bold=True)
        g = self.geometry
        bar = self._rect(slide, Box(4.45, 0.55, 0.04, g.height - 1.2), fill=t.accent)
        bar.shadow.inherit = False
        self._bullets(slide, Box(4.7, 0.5, g.width - 4.7 - g.margin_x, g.height - 1.1), spec.bullets, plan.body_pt)

    # ── 공통 요소 ────────────────────────────────────────────

    def _title(self, slide, text: str, pt: int) -> None:
        if not self._set_placeholder(slide, PP_PLACEHOLDER.TITLE, text, pt):
            self._text(slide, Box(0.5, 0.41, self.geometry.inner_width, 0.5), text, pt, self.theme.navy, bold=True)

    def _title_message(self, slide, spec, plan) -> None:
        self._title(slide, spec.title, plan.title_pt)
        if spec.message:
            self._text(slide, self.geometry.message_box(), spec.message, plan.message_pt, self.theme.navy,
                       bold=True, anchor=MSO_ANCHOR.MIDDLE)

    def _cards(self, slide, columns, has_message: bool, pt: int) -> None:
        boxes = column_boxes(self.geometry, has_message, len(columns))
        height = max(fitted_height(box, col.bullets, pt) for col, box in zip(columns, boxes))
        for col, box in zip(columns, boxes):
            self._card(slide, Box(box.left, box.top, box.width, height), col.heading, col.bullets, pt)

    def _card(self, slide, box: Box, heading: str, bullets: list[str], pt: int) -> None:
        t = self.theme
        self._rect(slide, box, fill=None, line=t.outline, rounded=True)
        header = self._rect(slide, Box(box.left, box.top, box.width, CARD_HEADER_H), fill=t.box)
        self._fill_text(header.text_frame, heading, 10, "FFFFFF", bold=True, align=PP_ALIGN.CENTER)
        inner = card_inner(box)
        tb = slide.shapes.add_textbox(Inches(inner.left), Inches(inner.top), Inches(inner.width), Inches(inner.height))
        self._bullets_in(tb.text_frame, bullets, pt, margin=0)

    def _footer(self, slide, number: int, refs: list[str]) -> None:
        t = self.theme
        for ph in slide.slide_layout.placeholders:
            if ph.placeholder_format.type == PP_PLACEHOLDER.FOOTER:
                el = copy.deepcopy(ph._element)
                el.nvSpPr.cNvPr.id = slide.shapes._next_shape_id  # 슬라이드 안에서 도형 ID가 겹치지 않게
                slide.shapes._spTree.append(el)
                footer = slide.shapes[-1]
                footer.text_frame.text = ""
                self._run(footer.text_frame.paragraphs[0], t.footer, None, None)
                break
        g = self.geometry
        self._text(slide, Box(g.width - 0.8, g.height - 0.32, 0.5, 0.24), str(number), 8, t.muted,
                   align=PP_ALIGN.RIGHT)
        if refs:
            label = "RFP 대응: " + ", ".join(refs)
            self._text(slide, Box(g.width - 5.3, g.height - 0.32, 4.45, 0.24), label, 7, t.muted,
                       align=PP_ALIGN.RIGHT)

    def _bullets(self, slide, box: Box, bullets: list[str], pt: int) -> None:
        tb = slide.shapes.add_textbox(Inches(box.left), Inches(box.top), Inches(box.width), Inches(box.height))
        self._bullets_in(tb.text_frame, bullets, pt, margin=0.05)

    def _bullets_in(self, tf, bullets: list[str], pt: int, margin: float) -> None:
        t = self.theme
        tf.word_wrap = True
        tf.auto_size = None
        tf.vertical_anchor = MSO_ANCHOR.TOP
        for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
            setattr(tf, side, Inches(margin))
        for i, raw in enumerate(bullets):
            text, level = split_bullet(raw)
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = PP_ALIGN.LEFT
            p.line_spacing = LINE_SPACING
            p.space_after = Pt(pt * 0.45)
            indent = BULLET_INDENT_IN + level * SUB_INDENT_IN
            _bullet(p, "•" if level == 0 else "–", indent, BULLET_INDENT_IN, t.navy if level == 0 else t.accent)
            self._rich(p, text, pt if level == 0 else max(pt - 1, 8), t.text)

    def _text(self, slide, box: Box, text: str, pt: int, color: str, bold=False, align=PP_ALIGN.LEFT,
              anchor=MSO_ANCHOR.TOP):
        tb = slide.shapes.add_textbox(Inches(box.left), Inches(box.top), Inches(box.width), Inches(box.height))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.auto_size = None
        tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = Inches(0.02)
        tf.margin_top = tf.margin_bottom = Inches(0.02)
        self._fill_text(tf, text, pt, color, bold=bold, align=align)
        return tb

    def _fill_text(self, tf, text: str, pt, color, bold=False, align=PP_ALIGN.LEFT) -> None:
        tf.word_wrap = True
        lines = text.split("\n")
        for i, line in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            self._rich(p, line, pt, color, bold=bold)

    def _rich(self, p, text: str, pt, color, bold=False) -> None:
        """'**강조**' 구간은 굵게 쓴다."""
        parts = text.split("**")
        for i, part in enumerate(parts):
            if part:
                self._run(p, part, pt, color, bold=bold or i % 2 == 1)

    def _run(self, p, text: str, pt, color, bold=False):
        r = p.add_run()
        r.text = text
        f = r.font
        if pt:
            f.size = Pt(pt)
        if bold:
            f.bold = True
        if color:
            f.color.rgb = _rgb(color)
        _set_typeface(r, self.theme.font)
        return r

    def _rect(self, slide, box: Box, fill: str | None, line: str | None = None, rounded=False, shape=None):
        shape_type = shape or (MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE)
        s = slide.shapes.add_shape(shape_type, Inches(box.left), Inches(box.top), Inches(box.width), Inches(box.height))
        s.shadow.inherit = False
        if rounded:
            s.adjustments[0] = 0.06
        if fill:
            s.fill.solid()
            s.fill.fore_color.rgb = _rgb(fill)
        else:
            s.fill.background()
        if line:
            s.line.color.rgb = _rgb(line)
            s.line.width = Pt(0.75)
            _line_alpha(s, 50)
        else:
            s.line.fill.background()
        tf = s.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE if fill else MSO_ANCHOR.TOP
        return s

    def _set_placeholder(self, slide, ph_type, text: str, pt: int) -> bool:
        for ph in slide.placeholders:
            if ph.placeholder_format.type == ph_type:
                tf = ph.text_frame
                tf.text = ""
                p = tf.paragraphs[0]
                for i, line in enumerate(text.split("\n")):
                    if i:
                        p = tf.add_paragraph()
                    self._run(p, line, pt, None)
                return True
        return False


def theme_from_template(prs) -> Theme:
    """템플릿의 테마 색과 예시 슬라이드 글꼴로 Theme을 만든다.

    Workday 팔레트면 예시 장표에서 뽑아 맞춘 기본값을 그대로 쓴다.
    """
    from collections import Counter

    from lxml import etree
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT

    theme_part = prs.slide_master.part.part_related_by(RT.THEME)
    root = etree.fromstring(theme_part.blob)
    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    scheme = root.find(".//a:clrScheme", ns)

    fonts: Counter[str] = Counter()
    for slide in prs.slides:
        for el in slide._element.iter(qn("a:ea"), qn("a:latin")):
            face = el.get("typeface", "")
            if face and not face.startswith("+"):
                fonts[face] += 1
    font = fonts.most_common(1)[0][0] if fonts else Theme.font

    if scheme is None or "workday" in (scheme.get("name") or "").lower():
        return Theme(font=font)

    def color(name: str, default: str) -> str:
        el = scheme.find(f"a:{name}/a:srgbClr", ns)
        if el is None:
            el = scheme.find(f"a:{name}/a:sysClr", ns)
            return (el.get("lastClr") if el is not None else None) or default
        return el.get("val", default)

    def dark(hex_: str) -> bool:
        r, g, b = (int(hex_[i:i + 2], 16) for i in (0, 2, 4))
        return 0.299 * r + 0.587 * g + 0.114 * b < 110

    dk1, dk2 = color("dk1", "000000"), color("dk2", "1F497D")
    navy = dk1 if dark(dk1) and dk1.upper() not in ("000000", "0D0D0D") else dk2
    accent = color("accent1", Theme.accent)
    return Theme(
        font=font,
        navy=navy,
        text=dk1 if dark(dk1) else "222222",
        box=dk2 if dark(dk2) else navy,
        accent=accent,
        accent_light=color("accent2", Theme.accent_light),
        outline=accent,
    )


# ── 부록: 요구사항 대응표 ─────────────────────────────────────


def with_appendix(deck: ProposalDeck, geo: Geometry) -> list:
    """덱 슬라이드 뒤(마무리 슬라이드 앞)에 요구사항 대응표 표 슬라이드를 붙인 목록."""
    slides = list(deck.slides)
    if not deck.requirements:
        return slides
    insert_at = len(slides) - 1 if slides and slides[-1].kind == "closing" else len(slides)

    columns = ["ID", "요구사항", "대응 유형", "대응 방안", "페이지"]
    ratios = [0.9, 3.0, 1.0, 3.2, 0.7]
    rows = []
    for req in deck.requirements.values():
        pages = deck.pages_for(req.id)
        rows.append([
            req.id,
            req.text,
            STATUS_LABELS.get(req.status or "", "미정"),
            req.response or "-",
            ", ".join(str(p) for p in pages) or "-",
        ])

    body = geo.body_box(False)
    widths = table_widths(ratios, len(columns), body.width)
    pt = TABLE_SIZES[-1]
    header_h = table_row_heights([columns], widths, pt)[0]
    pages: list[list[list[str]]] = [[]]
    used = header_h
    for row, h in zip(rows, table_row_heights(rows, widths, pt)):
        if pages[-1] and (used + h > body.height or len(pages[-1]) >= APPENDIX_ROWS_MAX):
            pages.append([])
            used = header_h
        pages[-1].append(row)
        used += h

    appendix = []
    for i, page_rows in enumerate(pages, 1):
        suffix = f" ({i}/{len(pages)})" if len(pages) > 1 else ""
        appendix.append(TableSlide(
            kind="table", title=APPENDIX_TITLE + suffix, columns=columns, rows=page_rows, column_widths=ratios,
        ))
    return slides[:insert_at] + appendix + slides[insert_at:]


# ── XML 도우미 ───────────────────────────────────────────────


def _clear_slides(prs) -> None:
    """템플릿에 들어 있는 예시 슬라이드를 모두 뺀다(레이아웃·마스터는 유지)."""
    sld_id_lst = prs.slides._sldIdLst
    for sld_id in list(sld_id_lst):
        prs.part.drop_rel(sld_id.rId)
        sld_id_lst.remove(sld_id)


def _find_layout(prs, role: str):
    by_name = {layout.name: layout for layout in prs.slide_layouts}
    for name in LAYOUT_CANDIDATES[role]:
        if name in by_name:
            return by_name[name]
    return prs.slide_layouts[0] if role == "cover" else prs.slide_layouts[-1]


def _drop_empty_placeholders(slide) -> None:
    for ph in list(slide.placeholders):
        if ph.placeholder_format.type == PP_PLACEHOLDER.FOOTER:
            continue
        if ph.has_text_frame and not ph.text_frame.text.strip():
            ph._element.getparent().remove(ph._element)


def _set_typeface(run, face: str) -> None:
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", face)


def _bullet(p, char: str, mar_left_in: float, hang_in: float, color: str) -> None:
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(Inches(mar_left_in)))
    pPr.set("indent", str(-Inches(hang_in)))
    for tag in ("a:buClr", "a:buSzPct", "a:buFont", "a:buNone", "a:buChar"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    bu_clr = pPr.makeelement(qn("a:buClr"), {})
    clr = bu_clr.makeelement(qn("a:srgbClr"), {"val": color})
    bu_clr.append(clr)
    bu_font = pPr.makeelement(qn("a:buFont"), {"typeface": "Arial"})
    bu_char = pPr.makeelement(qn("a:buChar"), {"char": char})
    # 순서: lnSpc, spcBef, spcAft 다음에 buClr → buFont → buChar
    anchor = None
    for tag in ("a:spcAft", "a:spcBef", "a:lnSpc"):
        anchor = pPr.find(qn(tag))
        if anchor is not None:
            break
    if anchor is None:
        pPr.insert(0, bu_char)
        pPr.insert(0, bu_font)
        pPr.insert(0, bu_clr)
    else:
        anchor.addnext(bu_char)
        anchor.addnext(bu_font)
        anchor.addnext(bu_clr)


def _line_alpha(shape, percent: int) -> None:
    ln = shape._element.spPr.find(qn("a:ln"))
    if ln is None:
        return
    clr = ln.find(qn("a:solidFill") + "/" + qn("a:srgbClr"))
    if clr is not None:
        alpha = clr.makeelement(qn("a:alpha"), {"val": str(percent * 1000)})
        clr.append(alpha)


def _plain_table_style(graphic_frame) -> None:
    """기본 표 스타일(줄무늬·강조 첫 행)을 끄고 직접 칠한 색만 쓰게 한다."""
    tbl = graphic_frame._element.graphic.graphicData.tbl
    tblPr = tbl.tblPr
    for attr in ("firstRow", "bandRow"):
        tblPr.set(attr, "0")
    style = tblPr.find(qn("a:tableStyleId"))
    if style is not None:
        tblPr.remove(style)


def _cell_borders(cell, color: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        for el in tcPr.findall(qn(tag)):
            tcPr.remove(el)
    fill = tcPr.find(qn("a:solidFill"))
    for i, tag in enumerate(("a:lnL", "a:lnR", "a:lnT", "a:lnB")):
        ln = tcPr.makeelement(qn(tag), {"w": str(Pt(0.75)), "cap": "flat", "cmpd": "sng"})
        solid = ln.makeelement(qn("a:solidFill"), {})
        solid.append(solid.makeelement(qn("a:srgbClr"), {"val": color}))
        ln.append(solid)
        # 테두리는 채우기(solidFill)보다 앞에 와야 한다.
        if fill is not None:
            fill.addprevious(ln)
        else:
            tcPr.append(ln)
