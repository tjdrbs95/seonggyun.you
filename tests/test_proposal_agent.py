"""모의 HTTP 응답으로 제안서 에이전트의 도구 실행 흐름 전체를 검증한다(API 키 불필요)."""

from __future__ import annotations

import csv
import json

import anthropic
import httpx2
from anthropic import DefaultHttpxClient
from pptx import Presentation

from proposal_agent.agent import ProposalAgent
from proposal_agent.cli import build_first_message
from proposal_agent.deck import ProposalDeck
from proposal_agent.render import DeckRenderer
from proposal_agent.tools import Workspace

from .conftest import ScriptedIO
from .proposal_fixtures import DEFAULT_TEMPLATE

LONG = "통합 인사 데이터 모델 기반으로 모든 계열사의 조직·인사·보상 정보를 단일 화면에서 관리하고 분석함" * 4


def _msg(i, content, stop="tool_use"):
    return {"id": f"msg_{i}", "type": "message", "role": "assistant", "model": "claude-opus-5",
            "content": content, "stop_reason": stop, "stop_sequence": None,
            "usage": {"input_tokens": 1, "output_tokens": 1}}


def _use(i, name, inp):
    return {"type": "tool_use", "id": f"toolu_{i}", "name": name, "input": inp}


SCRIPT = [
    _msg(1, [_use(1, "register_requirements", {"requirements": [
        {"id": "SFR-001", "category": "기능", "text": "그룹 통합 인사 정보 관리"},
        {"id": "SFR-002", "category": "기능", "text": "국내 급여 계산"},
    ]})]),
    _msg(2, [_use(2, "set_requirement_responses", {"responses": [
        {"id": "SFR-001", "status": "standard", "response": "Worker 모델로 통합"},
        {"id": "SFR-002", "status": "partner", "response": "국내 급여 파트너 솔루션 연동"},
    ]}), _use(3, "ask_user", {"questions": ["제안사명은?"]})]),
    _msg(3, [
        _use(4, "add_slide", {"slide": {"kind": "cover", "title": "○○ Workday 제안서", "subtitle": "메가존클라우드"}}),
        _use(5, "add_slide", {"slide": {"kind": "bullets", "title": "통합 인사", "message": "하나로 통합합니다.",
                                        "bullets": ["Worker 모델", "- 권한 분리"], "rfp_refs": ["SFR-001", "SFR-002"]}}),
        _use(6, "add_slide", {"slide": {"kind": "bullets", "title": "넘침", "bullets": [LONG] * 12}}),
        _use(7, "add_slide", {"slide": {"kind": "closing"}}),
    ]),
    _msg(4, [_use(8, "finalize", {})]),
    _msg(5, [{"type": "text", "text": "제안서를 완성했습니다."}], "end_turn"),
]


def test_agent_builds_pptx_and_matrix(tmp_path):
    requests = []
    responses = iter(SCRIPT)

    def handler(request):
        requests.append(request)
        return httpx2.Response(200, json=next(responses))

    client = anthropic.Anthropic(
        api_key="test", max_retries=0,
        http_client=DefaultHttpxClient(transport=httpx2.MockTransport(handler)),
    )
    io = ScriptedIO(["메가존클라우드"])
    ws = Workspace(ProposalDeck(), DeckRenderer(DEFAULT_TEMPLATE), tmp_path, io)
    agent = ProposalAgent(ws, client=client)

    rfp = {"type": "document", "source": {"type": "text", "media_type": "text/plain", "data": "SFR-001 통합"},
           "title": "rfp.md", "context": "RFP"}
    agent.send(build_first_message([rfp], "제안사: 메가존클라우드"))

    bodies = [json.loads(r.content) for r in requests]
    first_user = bodies[0]["messages"][0]["content"]
    assert first_user[0]["title"] == "rfp.md"
    assert "제안사: 메가존클라우드" in first_user[-1]["text"]
    add_slide_schema = next(t for t in bodies[0]["tools"] if t["name"] == "add_slide")["input_schema"]
    assert "discriminator" not in json.dumps(add_slide_schema)

    # 넘치는 슬라이드는 오류 메시지로 돌려주고 덱에는 넣지 않는다.
    results = {r["tool_use_id"]: r["content"] for r in bodies[3]["messages"][-1]["content"]}
    assert "넘칩니다" in results["toolu_6"]
    assert "2번 슬라이드" in results["toolu_5"]
    assert [s.kind for s in ws.deck.slides] == ["cover", "bullets", "closing"]
    assert ws.finalized

    prs = Presentation(str(ws.path()))
    assert len(prs.slides) == 4  # 표지, 본문, 부록 대응표, 마무리
    with ws.matrix_path().open(encoding="utf-8-sig") as f:
        rows = list(csv.reader(f))
    assert rows[1] == ["SFR-001", "기능", "그룹 통합 인사 정보 관리", "표준 기능", "Worker 모델로 통합", "2"]
    assert rows[2][3] == "파트너 솔루션"
    assert ws.state_path().exists()
    assert "제안서를 완성했습니다." in io.said
