import base64

import docx
import openpyxl
import pytest
from pptx import Presentation

from proposal_agent.inputs import InputError, build_blocks, collect_files, extract_text


def test_extracts_docx_xlsx_pptx_md(tmp_path):
    d = docx.Document()
    d.add_paragraph("1. 사업 개요")
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "SFR-001", "통합 인사"
    d.save(tmp_path / "rfp.docx")

    wb = openpyxl.Workbook()
    wb.active.title = "요구사항"
    wb.active.append(["ID", "내용"])
    wb.active.append(["SFR-002", "조직 개편"])
    wb.save(tmp_path / "req.xlsx")

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Workday Skills Cloud"
    prs.save(tmp_path / "sol.pptx")

    (tmp_path / "notes.md").write_text("# 회의록", encoding="utf-8")

    assert "SFR-001 | 통합 인사" in extract_text(tmp_path / "rfp.docx")
    assert "[시트 요구사항]" in extract_text(tmp_path / "req.xlsx")
    assert "[슬라이드 1]\nWorkday Skills Cloud" in extract_text(tmp_path / "sol.pptx")
    assert len(collect_files([tmp_path])) == 4


def make_pdf(path, pages):
    """글자가 들어 있는 최소 PDF를 만든다(Helvetica, ASCII 텍스트)."""
    objs = ["<< /Type /Catalog /Pages 2 0 R >>", None,
            "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    kids = []
    for text in pages:
        stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
        objs.append(f"<< /Length {len(stream)} >>\nstream\n{stream.decode()}\nendstream")
        content_id = len(objs)
        objs.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
                    f"/Resources << /Font << /F1 3 0 R >> >> /Contents {content_id} 0 R >>")
        kids.append(f"{len(objs)} 0 R")
    objs[1] = f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {len(kids)} >>"
    out = b"%PDF-1.4\n"
    offsets = []
    for i, body in enumerate(objs, 1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n{body}\nendobj\n".encode()
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    out += "".join(f"{o:010d} 00000 n \n" for o in offsets).encode()
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()
    path.write_bytes(out)


def test_solution_pdf_becomes_page_marked_text(tmp_path):
    pdf = tmp_path / "solution.pdf"
    make_pdf(pdf, ["Workday Payroll is available in the US and Canada.",
                   "Global Payroll Connect integrates partner payroll providers."])
    [block] = build_blocks([("Workday 솔루션 자료", [pdf], True)])
    assert block["source"]["type"] == "text"
    data = block["source"]["data"]
    assert data.startswith("[p.1]\nWorkday Payroll")
    assert "[p.2]\nGlobal Payroll Connect" in data


def test_scanned_pdf_falls_back_to_native(tmp_path):
    pdf = tmp_path / "scan.pdf"
    make_pdf(pdf, ["x", "y"])  # 페이지당 글자가 거의 없음 → 스캔 문서로 간주
    [block] = build_blocks([("Workday 솔루션 자료", [pdf], True)])
    assert block["source"]["type"] == "base64"


def test_pdf_becomes_base64_document(tmp_path):
    pdf = tmp_path / "rfp.pdf"
    pdf.write_bytes(b"%PDF-1.4 test")
    [block] = build_blocks([("RFP", [pdf], False)])
    assert block["source"]["media_type"] == "application/pdf"
    assert base64.b64decode(block["source"]["data"]) == b"%PDF-1.4 test"
    assert block["context"] == "RFP" and block["title"] == "rfp.pdf"


def test_rejects_hwp_and_missing(tmp_path):
    (tmp_path / "rfp.hwp").write_bytes(b"x")
    with pytest.raises(InputError, match="PDF로 변환"):
        collect_files([tmp_path / "rfp.hwp"])
    with pytest.raises(InputError, match="찾을 수 없습니다"):
        collect_files([tmp_path / "none.pdf"])
