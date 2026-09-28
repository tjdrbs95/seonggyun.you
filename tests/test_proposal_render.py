from pptx import Presentation

from proposal_agent.deck import SLIDE_ADAPTER, ProposalDeck, Requirement
from proposal_agent.render import APPENDIX_TITLE, DeckRenderer

from .proposal_fixtures import ALL_KINDS, DEFAULT_TEMPLATE


def _deck(n_requirements=3):
    deck = ProposalDeck()
    for d in ALL_KINDS:
        deck.add(SLIDE_ADAPTER.validate_python(d))
    deck.register([
        Requirement(id=f"SFR-{i:03d}", text=f"요구사항 {i} " + "설명 " * 10, status="standard", response="표준 기능 제공")
        for i in range(1, n_requirements + 1)
    ])
    return deck


def _texts(slide):
    out = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            out.append(shape.text_frame.text)
        if shape.has_table:
            out += [c.text for row in shape.table.rows for c in row.cells]
    return "\n".join(out)


def test_renders_every_kind_with_appendix_before_closing(tmp_path):
    out = tmp_path / "p.pptx"
    slides = DeckRenderer(DEFAULT_TEMPLATE).render(_deck(), out)
    prs = Presentation(str(out))

    assert len(prs.slides) == len(slides) == len(ALL_KINDS) + 1
    titles = [_texts(s) for s in prs.slides]
    assert "○○그룹 Workday HCM 구축 제안서" in titles[0]
    assert APPENDIX_TITLE in titles[-2]
    assert "SFR-001" in titles[-2] and "4, 5" in titles[-2]  # 대응 페이지
    assert "RFP 대응: SFR-001, SFR-002" in titles[4]
    assert prs.slides[4].notes_slide.notes_text_frame.text == "발표 스크립트"
    # 템플릿의 예시 슬라이드는 남지 않는다
    assert all("Click to" not in t for t in titles)


def test_long_requirement_list_paginates(tmp_path):
    out = tmp_path / "p.pptx"
    slides = DeckRenderer(DEFAULT_TEMPLATE).render(_deck(n_requirements=40), out)
    appendix = [s for s in slides if getattr(s, "title", "").startswith(APPENDIX_TITLE)]
    assert len(appendix) > 1
    assert appendix[0].title.endswith(f"(1/{len(appendix)})")
    assert sum(len(s.rows) for s in appendix) == 40


def test_korean_font_applied(tmp_path):
    out = tmp_path / "p.pptx"
    DeckRenderer(DEFAULT_TEMPLATE).render(_deck(), out)
    xml = Presentation(str(out)).slides[4]._element.xml
    assert 'typeface="Malgun Gothic"' in xml or "typeface=" in xml
